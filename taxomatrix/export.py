from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule


def export_matrix_csv(matrix: pd.DataFrame, output_path: Path) -> Path:
    matrix.to_csv(output_path, index=True, index_label="genome")
    return output_path


def export_label_mapping(label_mapping: pd.DataFrame, output_path: Path) -> Path:
    export_columns = [
        "original_filename",
        "assembly_accession",
        "organism_name",
        "strain",
        "display_label",
        "short_figure_label",
    ]
    label_mapping.loc[:, export_columns].to_csv(output_path, sep="\t", index=False)
    return output_path


def _auto_adjust_column_widths(worksheet) -> None:
    for column in worksheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column]
        width = max(len(value) for value in values) + 2
        worksheet.column_dimensions[column[0].column_letter].width = min(width, 40)


def _off_diagonal_values(matrix: pd.DataFrame) -> pd.Series:
    mask = pd.DataFrame(True, index=matrix.index, columns=matrix.columns)
    for label in matrix.index.intersection(matrix.columns):
        mask.loc[label, label] = False
    return matrix.where(mask).stack(future_stack=True).dropna()


def export_matrix_excel(
    matrix: pd.DataFrame,
    output_path: Path,
    metadata_sheet: pd.DataFrame | None = None,
) -> Path:
    summary_values = _off_diagonal_values(matrix)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        matrix.to_excel(writer, sheet_name="ANI Matrix", index=True, index_label="genome")
        summary = pd.DataFrame(
            {
                "metric": [
                    "genome_count",
                    "min_ani_excluding_diagonal",
                    "max_ani_excluding_diagonal",
                    "mean_ani_excluding_diagonal",
                ],
                "value": [
                    matrix.shape[0],
                    float(summary_values.min()) if not summary_values.empty else None,
                    float(summary_values.max()) if not summary_values.empty else None,
                    float(summary_values.mean()) if not summary_values.empty else None,
                ],
            }
        )
        summary.to_excel(writer, sheet_name="Summary", index=False)
        if metadata_sheet is not None:
            metadata_sheet.to_excel(writer, sheet_name="Metadata", index=False)

    workbook = load_workbook(output_path)
    worksheet = workbook["ANI Matrix"]
    worksheet.freeze_panes = "B2"

    max_row = worksheet.max_row
    max_col = worksheet.max_column
    for row in worksheet.iter_rows(min_row=2, min_col=2, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "0.00"

    if max_row >= 2 and max_col >= 2:
        worksheet.conditional_formatting.add(
            f"B2:{worksheet.cell(row=max_row, column=max_col).coordinate}",
            ColorScaleRule(
                start_type="num",
                start_value=70,
                start_color="D73027",
                mid_type="num",
                mid_value=85,
                mid_color="FEE08B",
                end_type="num",
                end_value=100,
                end_color="1A9850",
            ),
        )

    _auto_adjust_column_widths(worksheet)

    summary_sheet = workbook["Summary"]
    summary_sheet.freeze_panes = "A2"
    for row in summary_sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    _auto_adjust_column_widths(summary_sheet)

    if "Metadata" in workbook.sheetnames:
        metadata_worksheet = workbook["Metadata"]
        metadata_worksheet.freeze_panes = "A2"
        _auto_adjust_column_widths(metadata_worksheet)

    workbook.save(output_path)
    return output_path
